"""
Build outputs/Fraud_Detection_Dashboard.xlsx with openpyxl using REAL data
exported from the warehouse marts (data/processed/*.csv) and the fraud
model outputs (outputs/model_metrics.json, outputs/flagged_transactions.csv).

Sheets:
  Dashboard           - KPI cards + 4 charts
  Summary             - pre-aggregated tables the Dashboard sheet's charts reference
  RawData             - fact_transactions as an Excel Table (sampled if very large)
  TransactionFeatures - mart_transaction_features sample as an Excel Table
  FlaggedTransactions - flagged_transactions.csv as an Excel Table
"""
import json
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, BarChart, PieChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

PROCESSED_DIR = os.environ.get("PROCESSED_DIR", "/opt/airflow/data/processed")
OUTPUTS_DIR = os.environ.get("OUTPUTS_DIR", "/opt/airflow/outputs")
XLSX_PATH = os.path.join(OUTPUTS_DIR, "Fraud_Detection_Dashboard.xlsx")

HEADER_FILL = PatternFill(start_color="7A1F2B", end_color="7A1F2B", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
KPI_FILL = PatternFill(start_color="FCEEEF", end_color="FCEEEF", fill_type="solid")
KPI_LABEL_FONT = Font(size=10, color="7A1F2B", bold=True)
KPI_VALUE_FONT = Font(size=18, bold=True, color="7A1F2B")
THIN_BORDER = Border(*(Side(style="thin", color="CCCCCC"),) * 4)

MAX_RAWDATA_ROWS = 30000  # cap RawData sheet size for Excel performance; full data lives in Postgres


def add_table(ws, df, start_row=1, start_col=1, table_name="Table1"):
    for j, col in enumerate(df.columns):
        cell = ws.cell(row=start_row, column=start_col + j, value=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, val in enumerate(row):
            ws.cell(row=start_row + i, column=start_col + j, value=val)

    n_rows = len(df)
    n_cols = len(df.columns)
    last_col_letter = get_column_letter(start_col + n_cols - 1)
    first_col_letter = get_column_letter(start_col)
    ref = f"{first_col_letter}{start_row}:{last_col_letter}{start_row + n_rows}"
    tab = Table(displayName=table_name, ref=ref)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(tab)
    for j, col in enumerate(df.columns):
        col_letter = get_column_letter(start_col + j)
        max_len = max([len(str(col))] + [len(str(v)) for v in df[col].astype(str).head(200)])
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 40)
    return start_row + n_rows + 1


def main():
    os.makedirs(OUTPUTS_DIR, exist_ok=True)

    fact_txn = pd.read_csv(os.path.join(PROCESSED_DIR, "fact_transactions.csv"), parse_dates=["transaction_datetime"])
    features = pd.read_csv(os.path.join(PROCESSED_DIR, "transaction_features.csv"), parse_dates=["transaction_datetime"])
    flagged = pd.read_csv(os.path.join(OUTPUTS_DIR, "flagged_transactions.csv"))

    with open(os.path.join(OUTPUTS_DIR, "model_metrics.json")) as f:
        metrics = json.load(f)

    # ---- Pre-aggregations (real data) ----
    total_txns = len(fact_txn)
    total_fraud = int(fact_txn["is_fraud"].sum())
    fraud_rate = total_fraud / total_txns if total_txns else 0
    total_fraud_amount = float(fact_txn.loc[fact_txn["is_fraud"] == 1, "amount"].sum())
    avg_fraud_amount = float(fact_txn.loc[fact_txn["is_fraud"] == 1, "amount"].mean())
    model_auc = metrics["winner_metrics"]["auc_pr"]  # primary metric at ~1% fraud; AUC-ROC rounds to 1.000
    model_precision = metrics["winner_metrics"]["precision"]
    model_recall = metrics["winner_metrics"]["recall"]

    # fraud rate by merchant category
    cat_group = fact_txn.groupby("merchant_category").agg(
        total=("transaction_id", "count"), fraud=("is_fraud", "sum")
    )
    cat_group["fraud_rate_pct"] = (cat_group["fraud"] / cat_group["total"] * 100).round(3)
    cat_group = cat_group.sort_values("fraud_rate_pct", ascending=False).reset_index()
    cat_df = cat_group[["merchant_category", "fraud_rate_pct"]].copy()
    cat_df.columns = ["MerchantCategory", "FraudRatePct"]

    # fraud rate by hour of day
    hour_group = fact_txn.groupby("transaction_hour").agg(
        total=("transaction_id", "count"), fraud=("is_fraud", "sum")
    )
    hour_group["fraud_rate_pct"] = (hour_group["fraud"] / hour_group["total"] * 100).round(3)
    hour_group = hour_group.reset_index().sort_values("transaction_hour")
    hour_df = hour_group[["transaction_hour", "fraud_rate_pct"]].copy()
    hour_df.columns = ["Hour", "FraudRatePct"]

    # fraud amount distribution (bucketed)
    fraud_amounts = fact_txn.loc[fact_txn["is_fraud"] == 1, "amount"]
    bins = [0, 100, 250, 500, 1000, 2000, 100000]
    labels = ["$0-100", "$100-250", "$250-500", "$500-1000", "$1000-2000", "$2000+"]
    bucketed = pd.cut(fraud_amounts, bins=bins, labels=labels, right=False)
    amount_dist = bucketed.value_counts().reindex(labels).fillna(0).reset_index()
    amount_dist.columns = ["AmountRange", "FraudCount"]

    # top flagged high-risk transactions
    top_flagged = flagged.sort_values("fraud_probability", ascending=False).head(15)[
        ["transaction_id", "fraud_probability"]
    ].copy()
    top_flagged.columns = ["TransactionID", "FraudProbability"]

    # ---- Workbook ----
    wb = Workbook()

    ws_sum = wb.active
    ws_sum.title = "Summary"
    add_table(ws_sum, cat_df, start_row=1, start_col=1, table_name="TblCategoryFraudRate")
    add_table(ws_sum, hour_df, start_row=1, start_col=4, table_name="TblHourFraudRate")
    add_table(ws_sum, amount_dist, start_row=1, start_col=7, table_name="TblAmountDist")
    add_table(ws_sum, top_flagged, start_row=1, start_col=10, table_name="TblTopFlagged")

    # Dashboard sheet
    ws_dash = wb.create_sheet("Dashboard")
    ws_dash.sheet_view.showGridLines = False
    ws_dash["A1"] = "Bank Transaction Fraud Detection Dashboard"
    ws_dash["A1"].font = Font(size=20, bold=True, color="7A1F2B")
    ws_dash.merge_cells("A1:H1")

    kpis = [
        ("Total Transactions", f"{total_txns:,}"),
        ("Total Fraud Count", f"{total_fraud:,}"),
        ("Fraud Rate", f"{fraud_rate:.3%}"),
        ("Total Fraud $ Amount", f"${total_fraud_amount:,.0f}"),
        ("Avg Fraud Transaction", f"${avg_fraud_amount:,.2f}"),
        ("Model AUC-PR / Precision / Recall", f"{model_auc:.3f} / {model_precision:.3f} / {model_recall:.3f}"),
    ]
    for i, (label, value) in enumerate(kpis):
        col = 1 + (i % 3) * 3
        r = 3 + (i // 3) * 4
        cell_label = ws_dash.cell(row=r, column=col, value=label)
        cell_label.font = KPI_LABEL_FONT
        cell_value = ws_dash.cell(row=r + 1, column=col, value=value)
        cell_value.font = KPI_VALUE_FONT
        for rr in (r, r + 1):
            for cc in (col, col + 1):
                ws_dash.cell(row=rr, column=cc).fill = KPI_FILL
                ws_dash.cell(row=rr, column=cc).border = THIN_BORDER

    chart_anchor_row = 13

    # Chart 1: fraud rate by merchant category (bar)
    bar1 = BarChart()
    bar1.type = "bar"
    bar1.title = "Fraud Rate by Merchant Category (%)"
    bar1.style = 10
    data = Reference(ws_sum, min_col=2, min_row=1, max_row=1 + len(cat_df))
    cats = Reference(ws_sum, min_col=1, min_row=2, max_row=1 + len(cat_df))
    bar1.add_data(data, titles_from_data=True)
    bar1.set_categories(cats)
    bar1.width, bar1.height = 16, 9
    ws_dash.add_chart(bar1, f"A{chart_anchor_row}")

    # Chart 2: fraud rate by hour of day (line)
    line = LineChart()
    line.title = "Fraud Rate by Hour of Day (%)"
    line.style = 12
    line.y_axis.title = "Fraud Rate %"
    line.x_axis.title = "Hour"
    data = Reference(ws_sum, min_col=5, min_row=1, max_row=1 + len(hour_df))
    cats = Reference(ws_sum, min_col=4, min_row=2, max_row=1 + len(hour_df))
    line.add_data(data, titles_from_data=True)
    line.set_categories(cats)
    line.width, line.height = 16, 9
    ws_dash.add_chart(line, f"J{chart_anchor_row}")

    chart_anchor_row2 = chart_anchor_row + 19

    # Chart 3: fraud amount distribution (pie)
    pie = PieChart()
    pie.title = "Fraud Amount Distribution"
    data = Reference(ws_sum, min_col=8, min_row=1, max_row=1 + len(amount_dist))
    cats = Reference(ws_sum, min_col=7, min_row=2, max_row=1 + len(amount_dist))
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(cats)
    pie.width, pie.height = 16, 9
    ws_dash.add_chart(pie, f"A{chart_anchor_row2}")

    # Chart 4: top flagged high-risk transactions by probability (bar)
    bar2 = BarChart()
    bar2.type = "bar"
    bar2.title = "Top Flagged High-Risk Transactions (Fraud Probability)"
    bar2.style = 11
    data = Reference(ws_sum, min_col=11, min_row=1, max_row=1 + len(top_flagged))
    cats = Reference(ws_sum, min_col=10, min_row=2, max_row=1 + len(top_flagged))
    bar2.add_data(data, titles_from_data=True)
    bar2.set_categories(cats)
    bar2.width, bar2.height = 16, 9
    ws_dash.add_chart(bar2, f"J{chart_anchor_row2}")

    for col in "ABCDEFGHIJKLMNO":
        ws_dash.column_dimensions[col].width = 16

    # RawData sheet (fact_transactions) - capped for Excel performance
    ws_raw = wb.create_sheet("RawData")
    fact_display = fact_txn.copy()
    if len(fact_display) > MAX_RAWDATA_ROWS:
        fact_display = fact_display.sample(MAX_RAWDATA_ROWS, random_state=42).sort_values("transaction_datetime")
    fact_display["transaction_datetime"] = fact_display["transaction_datetime"].astype(str)
    add_table(ws_raw, fact_display, table_name="TblFactTransactions")

    # TransactionFeatures sheet - sample (real data, capped for performance)
    ws_feat = wb.create_sheet("TransactionFeatures")
    feat_display = features.copy()
    if len(feat_display) > MAX_RAWDATA_ROWS:
        feat_display = feat_display.sample(MAX_RAWDATA_ROWS, random_state=42).sort_values("transaction_datetime")
    feat_display["transaction_datetime"] = feat_display["transaction_datetime"].astype(str)
    add_table(ws_feat, feat_display, table_name="TblTransactionFeatures")

    # FlaggedTransactions sheet - full real data (flagged set is small)
    ws_flag = wb.create_sheet("FlaggedTransactions")
    add_table(ws_flag, flagged, table_name="TblFlaggedTransactions")

    wb.save(XLSX_PATH)
    print(f"Saved workbook: {XLSX_PATH}")
    print(f"Sheets: {wb.sheetnames}")

    return {
        "total_txns": total_txns,
        "total_fraud": total_fraud,
        "fraud_rate": fraud_rate,
        "total_fraud_amount": total_fraud_amount,
        "avg_fraud_amount": avg_fraud_amount,
        "model_auc": model_auc,
    }


if __name__ == "__main__":
    main()
